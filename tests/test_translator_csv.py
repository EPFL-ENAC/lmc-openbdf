import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from openbdf.cli.save import save_translator, translator_skeleton, workbook_from_csv_dir
from openbdf.lib.converters.translator_csv import csv_dir_rows, workbook_sheets_to_csv_dir
from openbdf.lib.converters.translators import ABOUT_SHEET_NAME, DocumentTranslator
from openbdf.main import app

TRANSLATOR_ROOT = Path(__file__).parents[1] / "translators"
TRANSLATOR_DIRS = {"translator_at": "full", "translator_bec_at": "bec", "translator_bmc_at": "bmc"}
runner = CliRunner()


def text_values(workbook) -> dict[str, dict[tuple[int, int], str]]:
    """Every non-empty cell of a workbook as text, keyed by sheet and coordinate."""
    return {
        sheet.title: {
            (cell.row, cell.column): str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        }
        for sheet in workbook.worksheets
    }


def test_csv_dir_round_trip_keeps_every_cell(tmp_path: Path):
    source = tmp_path / "translator.xlsx"
    save_translator(source)
    directory = tmp_path / "csv"
    workbook_sheets_to_csv_dir(load_workbook(source), directory, exclude=(ABOUT_SHEET_NAME,))

    rebuilt = load_workbook(workbook_from_csv_dir(directory))
    original = load_workbook(source)

    assert rebuilt.sheetnames == original.sheetnames
    assert text_values(rebuilt) == text_values(original)


def test_rebuilt_workbook_holds_the_sheets_of_the_directory_only(tmp_path: Path):
    source = tmp_path / "csv"
    workbook_sheets_to_csv_dir(load_workbook(translator_skeleton()), source, exclude=(ABOUT_SHEET_NAME,))
    subset = tmp_path / "subset"
    subset.mkdir()
    for name in ("002-General Information.csv", "015-ClimateZone.csv"):
        (subset / name).write_bytes((source / name).read_bytes())

    rebuilt = load_workbook(workbook_from_csv_dir(subset))

    assert rebuilt.sheetnames == [ABOUT_SHEET_NAME, "General Information", "ClimateZone"]


def test_rebuilt_sheet_takes_the_csv_header_and_keeps_the_template_style(tmp_path: Path):
    source = tmp_path / "csv"
    workbook_sheets_to_csv_dir(load_workbook(translator_skeleton("bec")), source, exclude=(ABOUT_SHEET_NAME,))
    path = source / "001-General Information.csv"
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream))
    rows[0][0] = "Custom header"
    with path.open("w", encoding="utf-8", newline="") as stream:
        csv.writer(stream, lineterminator="\n").writerows(rows)

    sheet = load_workbook(workbook_from_csv_dir(source))["General Information"]

    assert sheet["A1"].value == "Custom header"
    assert sheet["A1"].style == "AttributeName"
    assert sheet["A1"].font.bold
    assert sheet["A2"].style == "EditableValue"


def numeric_text(mapping: dict[str, list[str]]) -> dict[str, list[str]]:
    """The same mapping with one written form for numbers, because Excel reads 0 as 0.0."""

    def normalize(value: str) -> str:
        try:
            return str(float(value))
        except ValueError:
            return value

    return {key: [normalize(value) for value in values] for key, values in mapping.items()}


def test_from_csv_dir_matches_from_xlsx(tmp_path: Path):
    directory = TRANSLATOR_ROOT / "translator_at"
    rebuilt = tmp_path / "translator_at.xlsx"
    rebuilt.write_bytes(workbook_from_csv_dir(directory).getvalue())
    args = {"attribute_name_col": 0, "translated_name_col": 1, "default_value_col": 5, "preprocess_rule_col": 6}

    from_csv = DocumentTranslator.from_csv_dir(directory, **args)
    from_xlsx = DocumentTranslator.from_xlsx(rebuilt, **args)

    assert [item.model_dump() for item in from_csv._library._sheets_translations] == [
        item.model_dump() for item in from_xlsx._library._sheets_translations
    ]
    assert from_csv._translators.keys() == from_xlsx._translators.keys()
    for name, translator in from_xlsx._translators.items():
        assert numeric_text(from_csv._translators[name]._translation_dict) == numeric_text(translator._translation_dict)
        assert from_csv._translators[name]._default_values == translator._default_values
        assert from_csv._translators[name]._preprocess_rules == translator._preprocess_rules


@pytest.mark.parametrize("directory", TRANSLATOR_DIRS)
def test_committed_csv_dirs_cover_the_model(directory: str):
    rows_by_sheet = csv_dir_rows(TRANSLATOR_ROOT / directory)
    skeleton = load_workbook(translator_skeleton(TRANSLATOR_DIRS[directory]))

    assert list(rows_by_sheet) == [name for name in skeleton.sheetnames if name != ABOUT_SHEET_NAME]
    for sheet in skeleton.worksheets:
        assert len(sheet.title) <= 31
        if sheet.title == ABOUT_SHEET_NAME:
            continue
        model_keys = {str(sheet.cell(row=row, column=2).value) for row in range(2, sheet.max_row + 1)}
        csv_keys = {row[1] for row in rows_by_sheet[sheet.title][1:] if len(row) > 1}
        assert model_keys <= csv_keys


@pytest.mark.parametrize(
    "command",
    ["save-template", "save-translator-template", "save-bec-template", "save-bmc-template"],
)
def test_save_commands_write_a_csv_directory(command: str, tmp_path: Path):
    directory = tmp_path / command

    result = runner.invoke(app, [command, str(directory), "--csv"])

    assert result.exit_code == 0, result.output
    assert csv_dir_rows(directory)


def test_csv_dir_rejects_unknown_file_names(tmp_path: Path):
    (tmp_path / "notes.csv").write_text("", encoding="utf-8")

    with pytest.raises(ValueError, match="CSV file names"):
        csv_dir_rows(tmp_path)
