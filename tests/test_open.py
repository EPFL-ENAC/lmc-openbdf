from decimal import Decimal
from enum import Enum
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from typer.testing import CliRunner

import openbdf.main as main_module
from openbdf.cli.save import save_openbdf_to_xlsx
from openbdf.lib.types.annotations import unwrap_optional
from openbdf.main import app
from openbdf.model.building_bill_materials_slim import BuildingBillMaterialsSlimRecordBase
from openbdf.model.enums.materials import MaterialDataType
from openbdf.model.general_info import ProjectRecordBase
from openbdf.model.lca_results_full_wide import LCAResultsWideFullBase
from openbdf.model.project_material_database import ProjectMaterialDatabaseBase

runner = CliRunner()


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, dataframe in sheets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, header=False, index=False)


def default_openbdf_sheets() -> dict[str, pd.DataFrame]:
    project = pd.DataFrame([[None] * 11])
    project.iloc[0, 3] = "project_name"
    project.iloc[0, 9] = "Example project"

    # Generated template layout: field codes in row 2, data from row 9.
    bill_of_materials = pd.DataFrame([[None] * 2 for _ in range(9)])
    bill_of_materials.iloc[1, 1] = "material_id"
    bill_of_materials.iloc[8, 1] = "MAT-1"

    results_wide = pd.DataFrame([[None] * 2 for _ in range(9)])
    results_wide.iloc[1, 1] = "material_id"
    results_wide.iloc[8, 1] = "MAT-1"

    # Project Material Database in the repo wide layout: field codes in row 2, Reference row 8, data from row 9.
    project_material_database = pd.DataFrame([[None] * 2 for _ in range(9)])
    project_material_database.iloc[1, 1] = "material_ref_id"
    project_material_database.iloc[8, 1] = "M0001"

    return {
        "General Information": project,
        "Bill of Materials": bill_of_materials,
        "Project Material Database": project_material_database,
        "LCA Result (Wide)": results_wide,
    }


def mock_validation(monkeypatch) -> None:
    """This module tests the command wiring, so the fake converters skip the checks."""
    monkeypatch.setattr(main_module, "validate_project", lambda *args, **kwargs: [])


def mock_conversion_and_save(monkeypatch, captured_tables: list[Any]) -> None:
    def convert(tables, document_translator, *, findings=None, draft=False, drop_invalid=False):
        captured_tables.append((tables, drop_invalid))
        return object()

    monkeypatch.setattr(main_module, "project_package_from_tables", convert)
    monkeypatch.setattr(main_module, "save_openbdf_to_xlsx", lambda **kwargs: None)
    mock_validation(monkeypatch)


def test_open_command_processes_sheets_split_across_files(tmp_path: Path, monkeypatch):
    sheets = default_openbdf_sheets()
    project = tmp_path / "project.xlsx"
    materials = tmp_path / "materials.xlsx"
    materials_database = tmp_path / "materials_database.xlsx"
    results = tmp_path / "results.xlsx"
    write_workbook(project, {"General Information": sheets["General Information"]})
    write_workbook(materials, {"Bill of Materials": sheets["Bill of Materials"]})
    write_workbook(materials_database, {"Project Material Database": sheets["Project Material Database"]})
    write_workbook(results, {"LCA Result (Wide)": sheets["LCA Result (Wide)"]})
    captured_tables: list[Any] = []
    mock_conversion_and_save(monkeypatch, captured_tables)

    result = runner.invoke(app, ["open", str(results), str(project), str(materials), str(materials_database)])

    assert result.exit_code == 0
    assert len(captured_tables) == 1
    tables, drop_invalid = captured_tables[0]
    assert drop_invalid is False
    assert tables.project.to_dict(orient="records") == [{"field_code": "project_name", "value": "Example project"}]
    assert tables.bill_of_materials.to_dict(orient="records") == [{"material_id": "MAT-1"}]
    # start_row=8 skips the metadata rows, so the first record is the first data row.
    raw_project_material_database = tables.project_material_database.to_dict(orient="records")
    assert raw_project_material_database == [{"material_ref_id": "M0001"}]
    assert tables.results_wide.to_dict(orient="records") == [{"material_id": "MAT-1"}]


def test_open_command_accepts_one_file(tmp_path: Path, monkeypatch):
    workbook = tmp_path / "complete.xlsx"
    write_workbook(workbook, default_openbdf_sheets())
    captured_tables: list[Any] = []
    mock_conversion_and_save(monkeypatch, captured_tables)

    result = runner.invoke(app, ["open", str(workbook)])

    assert result.exit_code == 0
    assert len(captured_tables) == 1


def test_open_command_uses_draft_conversion(tmp_path: Path, monkeypatch):
    workbook = tmp_path / "complete.xlsx"
    write_workbook(workbook, default_openbdf_sheets())
    captured: list[Any] = []

    def convert(tables, document_translator, *, findings=None, draft=False, drop_invalid=False):
        captured.append(draft)
        return object()

    monkeypatch.setattr(main_module, "project_package_from_tables", convert)
    monkeypatch.setattr(main_module, "save_openbdf_to_xlsx", lambda **kwargs: None)
    mock_validation(monkeypatch)

    result = runner.invoke(app, ["open", str(workbook), "--draft"])

    assert result.exit_code == 0
    assert captured == [True]


def test_open_command_drop_invalid_materials_passes_flag(tmp_path: Path, monkeypatch):
    workbook = tmp_path / "complete.xlsx"
    write_workbook(workbook, default_openbdf_sheets())
    captured: list[Any] = []

    def convert_strict(tables, document_translator, *, findings=None, draft=False, drop_invalid=False):
        captured.append(drop_invalid)
        return object()

    monkeypatch.setattr(main_module, "project_package_from_tables", convert_strict)
    monkeypatch.setattr(main_module, "save_openbdf_to_xlsx", lambda **kwargs: None)
    mock_validation(monkeypatch)

    result = runner.invoke(app, ["open", str(workbook), "--drop-invalid-materials"])

    assert result.exit_code == 0
    assert captured == [True]


def test_open_command_drop_invalid_materials_default_is_false(tmp_path: Path, monkeypatch):
    workbook = tmp_path / "complete.xlsx"
    write_workbook(workbook, default_openbdf_sheets())
    captured: list[Any] = []

    def convert_strict(tables, document_translator, *, findings=None, draft=False, drop_invalid=False):
        captured.append(drop_invalid)
        return object()

    monkeypatch.setattr(main_module, "project_package_from_tables", convert_strict)
    monkeypatch.setattr(main_module, "save_openbdf_to_xlsx", lambda **kwargs: None)
    mock_validation(monkeypatch)

    result = runner.invoke(app, ["open", str(workbook)])

    assert result.exit_code == 0
    assert captured == [False]


def test_open_command_reports_duplicate_sheets(tmp_path: Path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    write_workbook(first, {"Shared": pd.DataFrame([[1]])})
    write_workbook(second, {"Shared": pd.DataFrame([[2]])})

    result = runner.invoke(app, ["open", str(first), str(second)])

    assert result.exit_code == 1
    assert "Duplicate sheet names" in result.output
    assert "Shared" in result.output
    assert str(first) in result.output
    assert str(second) in result.output


def test_open_command_uses_translator_option(tmp_path: Path, monkeypatch):
    workbook = tmp_path / "complete.xlsx"
    translator_path = tmp_path / "translator.xlsx"
    write_workbook(workbook, default_openbdf_sheets())
    captured_paths: list[str] = []
    captured_tables: list[Any] = []
    default_translator = main_module.DocumentTranslator.default_openbdf_translator()

    def load_translator(*, xlsx_path, **kwargs):
        captured_paths.append(xlsx_path)
        return default_translator

    monkeypatch.setattr(main_module.DocumentTranslator, "from_xlsx", load_translator)
    mock_conversion_and_save(monkeypatch, captured_tables)

    result = runner.invoke(app, ["open", str(workbook), "--translator", str(translator_path)])

    assert result.exit_code == 0
    assert captured_paths == [str(translator_path)]
    assert len(captured_tables) == 1


def test_open_command_merges_multiple_translators(tmp_path: Path, monkeypatch):
    workbook = tmp_path / "complete.xlsx"
    write_workbook(workbook, default_openbdf_sheets())
    translator_a = tmp_path / "translator_a.xlsx"
    translator_b = tmp_path / "translator_b.xlsx"

    from openbdf.lib.converters.translators import (
        CategoryTranslator,
        DocumentTranslator,
        TranslationsLibrary,
    )

    def make_translator(category: str):
        extra = DocumentTranslator(
            library=TranslationsLibrary([]),
            translators={
                category: CategoryTranslator({f"src_{category}": [f"dst_{category}"]}),
            },
        )
        return DocumentTranslator.default_openbdf_translator() + extra

    translator_a_obj = make_translator("General Information")
    translator_b_obj = make_translator("Bill of Materials")

    def load_translator(*, xlsx_path, **kwargs):
        return translator_a_obj if xlsx_path == str(translator_a) else translator_b_obj

    monkeypatch.setattr(main_module.DocumentTranslator, "from_xlsx", load_translator)
    captured_translators: list[Any] = []

    def convert(tables, document_translator, *, findings=None, draft=False, drop_invalid=False):
        captured_translators.append(document_translator)
        return object()

    monkeypatch.setattr(main_module, "project_package_from_tables", convert)
    monkeypatch.setattr(main_module, "save_openbdf_to_xlsx", lambda **kwargs: None)
    mock_validation(monkeypatch)

    result = runner.invoke(
        app,
        ["open", str(workbook), "-t", str(translator_a), "-t", str(translator_b)],
    )

    assert result.exit_code == 0
    assert len(captured_translators) == 1
    merged = captured_translators[0]
    assert merged._translators["General Information"] == translator_a_obj._translators["General Information"]
    assert merged._translators["Bill of Materials"] == translator_b_obj._translators["Bill of Materials"]


def valid_template_value(field_info: Any) -> Any:
    annotation, _ = unwrap_optional(field_info.annotation)
    if isinstance(annotation, type) and issubclass(annotation, Enum):
        return next(iter(annotation)).value
    if annotation is bool:
        return True
    if annotation in (int, Decimal):
        return 1
    return "test"


def fill_required_vertical(sheet, model_cls) -> None:
    """Fill required fields into a generated template's vertical sheet."""
    field_code_rows = {sheet.cell(row=row, column=4).value: row for row in range(2, sheet.max_row + 1)}
    for name, info in model_cls.model_fields.items():
        if info.is_required() and name in field_code_rows:
            sheet.cell(row=field_code_rows[name], column=10, value=valid_template_value(info))


def fill_required_horizontal(sheet, model_cls) -> None:
    """Fill required fields into the first data row of a generated template's horizontal sheet."""
    field_code_cols = {cell.value: cell.column for cell in sheet[2] if cell.value}
    for name, info in model_cls.model_fields.items():
        if info.is_required() and name in field_code_cols:
            sheet.cell(row=9, column=field_code_cols[name], value=valid_template_value(info))


def test_open_command_opens_filled_generated_template_as_single_file(tmp_path: Path):
    template_path = tmp_path / "template.xlsx"
    save_openbdf_to_xlsx(template_path, None)

    workbook = load_workbook(template_path)
    fill_required_vertical(workbook["General Information"], ProjectRecordBase)

    bom = workbook["Bill of Materials"]
    fill_required_horizontal(bom, BuildingBillMaterialsSlimRecordBase)
    bom_codes = {cell.value: cell.column for cell in bom[2] if cell.value}
    bom.cell(row=9, column=bom_codes["material_id"], value="MAT-1")
    bom.cell(row=9, column=bom_codes["custom_material_name"], value="Concrete C30/37")
    bom.cell(row=9, column=bom_codes["material_reference_service_life"], value=50)

    project_material_database = workbook["Project Material Database"]
    fill_required_horizontal(project_material_database, ProjectMaterialDatabaseBase)
    pmd_codes = {cell.value: cell.column for cell in project_material_database[2] if cell.value}
    project_material_database.cell(row=9, column=pmd_codes["material_ref_id"], value="M0001")
    project_material_database.cell(row=9, column=pmd_codes["custom_material_name"], value="Concrete C30/37")
    # A material is the pair of its name and its EPD dataset name, so both sheets must agree on it.
    bom.cell(
        row=9,
        column=bom_codes["lca_dataset_material_name"],
        value=project_material_database.cell(row=9, column=pmd_codes["lca_dataset_material_name"]).value,
    )

    wide = workbook["LCA Result (Wide)"]
    fill_required_horizontal(wide, LCAResultsWideFullBase)
    wide_codes = {cell.value: cell.column for cell in wide[2] if cell.value}
    wide.cell(row=9, column=wide_codes["material_id"], value="MAT-1")
    wide.cell(row=9, column=wide_codes["gwp_total_a1_a3"], value=123.45)
    workbook.save(template_path)

    output_path = tmp_path / "out.xlsx"
    result = runner.invoke(app, ["open", str(template_path), "-o", str(output_path)])

    assert result.exit_code == 0
    out = load_workbook(output_path)
    out_bom_codes = {cell.value: cell.column for cell in out["Bill of Materials"][2] if cell.value}
    out_bom = out["Bill of Materials"]
    assert out_bom.cell(row=9, column=out_bom_codes["material_id"]).value == "MAT-1"
    assert out_bom.cell(row=9, column=out_bom_codes["custom_material_name"]).value == "Concrete C30/37"
    assert out_bom.cell(row=9, column=out_bom_codes["material_reference_service_life"]).value == "50"
    out_wide_codes = {cell.value: cell.column for cell in out["LCA Result (Wide)"][2] if cell.value}
    out_wide = out["LCA Result (Wide)"]
    assert out_wide.cell(row=9, column=out_wide_codes["material_id"]).value == "MAT-1"
    assert out_wide.cell(row=9, column=out_wide_codes["gwp_total_a1_a3"]).value == "123.45"
    out_pmd_codes = {cell.value: cell.column for cell in out["Project Material Database"][2] if cell.value}
    out_pmd = out["Project Material Database"]
    assert out_pmd.cell(row=9, column=out_pmd_codes["material_ref_id"]).value == "M0001"
    assert out_pmd.cell(row=9, column=out_pmd_codes["custom_material_name"]).value == "Concrete C30/37"


def test_document_translator_add_last_wins_on_category_conflict():
    from openbdf.lib.converters.translators import (
        CategoryTranslator,
        DocumentTranslator,
        TranslationsLibrary,
    )

    first = DocumentTranslator(
        library=TranslationsLibrary([]),
        translators={"cat": CategoryTranslator({"src": ["first"]})},
    )
    second = DocumentTranslator(
        library=TranslationsLibrary([]),
        translators={"cat": CategoryTranslator({"src": ["second"]})},
    )

    merged = first + second
    assert merged._translators["cat"].translate("src") == ["second"]


def test_open_command_collects_every_finding_of_a_workbook_with_bad_cells(tmp_path: Path):
    template_path = tmp_path / "template.xlsx"
    save_openbdf_to_xlsx(template_path, None)

    workbook = load_workbook(template_path)
    project_sheet = workbook["General Information"]
    fill_required_vertical(project_sheet, ProjectRecordBase)
    project_code_rows = {cell.value: cell.row for cell in project_sheet["D"] if cell.value}
    project_sheet.cell(row=project_code_rows["building_gross_floor_area"], column=10, value=0)

    project_material_database = workbook["Project Material Database"]
    fill_required_horizontal(project_material_database, ProjectMaterialDatabaseBase)
    database_codes = {cell.value: cell.column for cell in project_material_database[2] if cell.value}
    project_material_database.cell(row=9, column=database_codes["material_ref_id"], value="M0001")
    project_material_database.cell(row=9, column=database_codes["custom_material_name"], value="Concrete C30/37")
    dataset_name = project_material_database.cell(row=9, column=database_codes["lca_dataset_material_name"]).value

    bom = workbook["Bill of Materials"]
    fill_required_horizontal(bom, BuildingBillMaterialsSlimRecordBase)
    bom_codes = {cell.value: cell.column for cell in bom[2] if cell.value}
    bom.cell(row=9, column=bom_codes["material_id"], value="MAT-1")
    bom.cell(row=9, column=bom_codes["custom_material_name"], value="Concrete C30/37")
    bom.cell(row=9, column=bom_codes["lca_dataset_material_name"], value=dataset_name)
    bom.cell(row=9, column=bom_codes["material_data_type"], value="bogus")
    for column in range(1, bom.max_column + 1):
        bom.cell(row=10, column=column, value=bom.cell(row=9, column=column).value)
    bom.cell(row=10, column=bom_codes["material_weight_kg"], value=100)
    bom.cell(row=10, column=bom_codes["material_data_type"], value=MaterialDataType.CPR_DOPC.value)

    wide = workbook["LCA Result (Wide)"]
    fill_required_horizontal(wide, LCAResultsWideFullBase)
    wide_codes = {cell.value: cell.column for cell in wide[2] if cell.value}
    wide.cell(row=9, column=wide_codes["material_id"], value="MAT-2")
    workbook.save(template_path)

    output_path = tmp_path / "out.xlsx"
    result = runner.invoke(app, ["open", str(template_path), "-o", str(output_path)])

    assert result.exit_code == 1
    floor_area_row = project_code_rows["building_gross_floor_area"]
    assert result.output.splitlines() == [
        "General Information:",
        f"  Row {floor_area_row}: E17: Floor area must be greater than zero",
        "Bill of Materials:",
        "  E32: material_id 'MAT-1' is used by 2 bill of materials rows",
        "  Row 9: E02: Value 'bogus' is not a valid MaterialDataType",
        "  E16: Row 9 has no mass, material efficiency analysis is not possible",
        "LCA Result (Wide):",
        "  E05: Bill of Materials has 2 rows, LCA Result (Wide) has 1",
        "  E06: Row 9: material_id MAT-1 in the inventory, MAT-2 in the results",
        "The workbook has reject findings, so it was not saved. Use --draft to save it anyway.",
    ]
    assert not output_path.exists()
