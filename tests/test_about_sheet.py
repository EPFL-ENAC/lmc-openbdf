import pandas as pd
from openpyxl import load_workbook

from openbdf.cli.save import ABOUT_SHEET_NAME, about_to_xlsx, register_named_styles


def test_about_sheet_renders_links_and_strips_markdown(tmp_path):
    output_path = tmp_path / "about.xlsx"
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        register_named_styles(writer)
        about_to_xlsx(writer)

    sheet = load_workbook(output_path)[ABOUT_SHEET_NAME]
    assert sheet.max_column == 1

    links = [cell for row in sheet.iter_rows(max_col=1) for cell in row if cell.hyperlink]
    assert [cell.value for cell in links] == [
        "nicole.disanto@epfl.ch",
        "martin.roeck@riselabs.at",
        "https://github.com/EPFL-ENAC/lmc-openbdf/",
    ]
    assert [cell.hyperlink.target for cell in links] == [
        "mailto:nicole.disanto@epfl.ch",
        "mailto:martin.roeck@riselabs.at",
        "https://github.com/EPFL-ENAC/lmc-openbdf/",
    ]

    for row in sheet.iter_rows(max_col=1):
        for cell in row:
            value = str(cell.value)
            assert len(value) <= 80
            assert "**" not in value
            assert not value.startswith(("#", ">"))
    assert sheet.column_dimensions["A"].width == 85
