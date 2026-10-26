from openpyxl import load_workbook

from openbdf.cli.save import save_openbdf_to_xlsx, save_translator
from openbdf.lib.converters.translators import (
    COMMON_CATEGORY_PROPERTIES,
    SHEETS_PROPERTIES_HEADERS,
)

FIELD_SHEET_NAMES = [properties["name"] for properties in COMMON_CATEGORY_PROPERTIES.values()]
FIELD_SHEET_HEADERS = [
    "Field code in your format",
    "OpenBDF field code",
    "OpenBDF attribute name",
    "Datatype",
    "Requirements",
    "Default value",
    "Preprocess rule",
]
MARKED_REQUIREMENTS = {"Required", "Strongly recommended", "Automatic"}
ENUM_SHEET_HEADERS = [
    "Value in your format",
    "OpenBDF enum value",
    "OpenBDF value display name",
    "OpenBDF value description (if any)",
    "OpenBDF value parent category (if applicable)",
]


def test_translator_template_adds_default_and_preprocess_columns(tmp_path):
    output_path = tmp_path / "translator_template.xlsx"
    save_translator(output_path)

    workbook = load_workbook(output_path)
    for sheet_name in FIELD_SHEET_NAMES:
        sheet = workbook[sheet_name]

        assert [sheet.cell(row=1, column=column).value for column in range(1, 8)] == FIELD_SHEET_HEADERS
        assert {sheet.cell(row=1, column=column).style for column in range(1, 8)} == {"AttributeName"}
        assert sheet.column_dimensions["F"].width == 30
        assert sheet.column_dimensions["G"].width == 30

        for row in range(2, sheet.max_row + 1):
            assert sheet.cell(row=row, column=6).value is None
            assert sheet.cell(row=row, column=7).value is None
            assert sheet.cell(row=row, column=6).style == "EditableValue"
            assert sheet.cell(row=row, column=7).style == "EditableValue"
            assert sheet.cell(row=row, column=6).protection.locked is False
            assert sheet.cell(row=row, column=7).protection.locked is False
            assert sheet.cell(row=row, column=1).style == "EditableValue"

    library = workbook["Translations Library"]
    assert library["B2"].style == "Normal"
    assert {
        library.cell(row=row, column=column).style
        for row in range(2, library.max_row + 1)
        for column in (1, 3, 4, 5, 10)
    } == {"EditableValue"}

    enum_sheet = workbook["ProjectUnits"]
    enum_rows = range(2, enum_sheet.max_row + 1)
    assert {enum_sheet.cell(row=row, column=1).style for row in enum_rows} == {"EditableValue"}

    for sheet in workbook.worksheets:
        if sheet.title == "About":
            continue
        assert {cell.style for cell in sheet[1] if cell.value is not None} == {"AttributeName"}

    bill_of_materials = workbook[COMMON_CATEGORY_PROPERTIES["bom"]["name"]]
    service_life_row = next(
        row for row in bill_of_materials.iter_rows() if row[1].value == "material_reference_service_life"
    )
    assert [cell.value for cell in service_life_row[:5]] == [
        None,
        "material_reference_service_life",
        "Reference Service Life (RSL)",
        "Decimal",
        "Strongly recommended",
    ]


def test_openbdf_template_requirements_and_reference_layout(tmp_path):
    output_path = tmp_path / "openbdf_template.xlsx"
    save_openbdf_to_xlsx(output_path, None)

    workbook = load_workbook(output_path)
    for sheet in workbook.worksheets:
        assert sheet.row_dimensions[1].height == 38
        for cell in sheet[1]:
            assert (cell.alignment.horizontal, cell.alignment.vertical) == ("left", "bottom")
    info = workbook["General Information"]
    assert [info.cell(row=1, column=column).value for column in range(9, 12)] == [
        "Requirements",
        "Value",
        "Reference",
    ]
    assert info.cell(row=2, column=9).value == "Automatic"
    assert info.cell(row=2, column=11).value == "ECHO; GBDB"

    # An automatic value is grey, so a contributor sees which values the tool computes.
    automatic_rows = [cell.row for cell in info["I"] if cell.value == "Automatic"]
    editable_rows = [cell.row for cell in info["I"] if cell.row > 1 and cell.value not in ("Automatic", None)]
    assert automatic_rows and editable_rows
    assert {info.cell(row=row, column=10).style for row in automatic_rows} == {"AutomaticValue"}
    assert {info.cell(row=row, column=10).style for row in editable_rows} == {"EditableValue"}
    assert {info.cell(row=row, column=column).style for row in range(2, info.max_row + 1) for column in (7, 8, 11)} == {
        "Metadata"
    }
    # The marked levels paint the five metadata columns of a General Information row.
    for row in range(2, info.max_row + 1):
        requirement = info.cell(row=row, column=9).value
        expected = requirement if requirement in MARKED_REQUIREMENTS else "Normal"
        assert {info.cell(row=row, column=column).style for column in range(1, 6)} == {expected}

    first_field_requirements = {
        "Bill of Materials": "Automatic",
        "Project Material Database": "Automatic",
        "LCA Result (Wide)": "Required",
        "LCA Result (Tall)": "Required",
    }
    reference_labels = {
        "Bill of Materials": "Reference",
        "Project Material Database": "Reference",
        "LCA Result (Wide)": "Reference Alignment",
        "LCA Result (Tall)": "Reference Alignment",
    }

    for sheet_name in FIELD_SHEET_NAMES[1:]:
        sheet = workbook[sheet_name]
        assert sheet["A2"].value == "Field Code"
        assert sheet["A3"].value == "Description"
        assert sheet["A4"].value == "Data Type"
        assert sheet["A5"].value == "Constraint"
        assert sheet["A6"].value == "Units"
        assert sheet["A7"].value == "Requirements"
        assert sheet["A8"].value == reference_labels[sheet_name]
        # The reference row keeps the first data row at row 9 on every field sheet.
        assert sheet["A9"].value is None
        for cell in sheet[7]:
            if cell.column > 1 and cell.value:
                expected = "AutomaticValue" if cell.value == "Automatic" else "EditableValue"
                assert sheet.cell(row=9, column=cell.column).style == expected
        # Units and constraints are grey, whether they carry a value or not.
        for row in (5, 6):
            assert {sheet.cell(row=row, column=cell.column).style for cell in sheet[row] if cell.column > 1} == {
                "Metadata"
            }
        # The first column labels every header row with white bold text on green.
        assert {sheet.cell(row=row, column=1).style for row in range(1, 9)} == {"AttributeName"}
        # Description and reference cells carry the colour of the marked levels.
        for cell in sheet[7]:
            if cell.column == 1 or not cell.value:
                continue
            expected = cell.value if cell.value in MARKED_REQUIREMENTS else "Normal"
            assert sheet.cell(row=3, column=cell.column).style == expected
            expected = cell.value if cell.value in MARKED_REQUIREMENTS else "Metadata"
            assert sheet.cell(row=8, column=cell.column).style == expected
        assert sheet.cell(row=2, column=2).value == (
            "material_ref_id" if sheet_name == "Project Material Database" else "material_id"
        )
        assert sheet.cell(row=7, column=2).value == first_field_requirements[sheet_name]

    type_definitions = workbook["TypeDefinitions"]
    assert type_definitions["A1"].style == "AttributeName"
    assert type_definitions["B1"].style == "AttributeName"
    assert type_definitions["A2"].value == "Reference"
    assert type_definitions["B2"].value == "GBDI; SLiCe"
    assert {cell.style for cell in type_definitions[2] if cell.column > 1} == {"Metadata"}
    # Enum values start at row 3, because row 2 holds the datatype reference, and are read-only metadata.
    body = [cell for row in type_definitions.iter_rows(min_row=3) for cell in row if cell.value is not None]
    assert body
    assert {cell.style for cell in body} == {"Metadata"}
    assert type_definitions["B3"].value is not None

    theme = workbook.loaded_theme
    assert b'<a:latin typeface="Calibri"/>' in theme
    assert b"Cambria" not in theme

    assert all(font.name == "Calibri" for font in workbook._fonts)


def test_translator_template_keeps_non_field_sheet_layouts(tmp_path):
    output_path = tmp_path / "translator_template.xlsx"
    save_translator(output_path)

    workbook = load_workbook(output_path, read_only=True)
    library = workbook["Translations Library"]
    expected_library_headers = list(SHEETS_PROPERTIES_HEADERS.values())

    assert library.max_column == len(expected_library_headers)
    assert [library.cell(row=1, column=column).value for column in range(1, 11)] == expected_library_headers

    enum_sheet = workbook["ProjectUnits"]
    assert enum_sheet.max_column == 5
    assert [enum_sheet.cell(row=1, column=column).value for column in range(1, 6)] == ENUM_SHEET_HEADERS
